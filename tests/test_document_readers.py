"""Document text extraction tests without external services or model downloads."""

from __future__ import annotations

from pathlib import Path

import pytest
import rag_mcp.document_readers as document_readers
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from rag_mcp.chunking import chunk_document, discover_documents
from rag_mcp.config import RagConfig
from rag_mcp.document_readers import DocumentReadError, read_document_text
from rag_mcp.service_knowledge import DenseDocumentIngester, KnowledgeValidationError


def test_pdf_text_is_extracted_and_chunked(tmp_path: Path) -> None:
    source = tmp_path / "manual.pdf"
    _write_text_pdf(source, "Galaxy X7 E03 reset procedure")

    chunks = chunk_document(source, source_root=tmp_path)

    assert len(chunks) == 1
    assert chunks[0].source == "manual.pdf"
    assert chunks[0].title == "Galaxy X7 E03 reset procedure"
    assert "E03 reset" in chunks[0].text
    assert chunks[0].metadata["document_suffix"] == ".pdf"


def test_docx_paragraphs_and_tables_are_extracted(tmp_path: Path) -> None:
    source = tmp_path / "guide.docx"
    document = Document()
    document.add_heading("星河 X7 故障手册", level=1)
    document.add_paragraph("E03 表示温度传感器异常。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "错误码"
    table.cell(0, 1).text = "处理方法"
    table.cell(1, 0).text = "E03"
    table.cell(1, 1).text = "检查传感器连接"
    document.save(source)

    text = read_document_text(source)
    chunks = chunk_document(source, source_root=tmp_path)

    assert "星河 X7 故障手册" in text
    assert "E03 表示温度传感器异常" in text
    assert "错误码 | 处理方法" in text
    assert "E03 | 检查传感器连接" in text
    assert chunks[0].metadata["document_suffix"] == ".docx"


def test_blank_pdf_reports_that_ocr_is_required(tmp_path: Path) -> None:
    source = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with source.open("wb") as handle:
        writer.write(handle)

    with pytest.raises(DocumentReadError, match="OCR"):
        read_document_text(source)


def test_password_protected_pdf_has_safe_error(tmp_path: Path) -> None:
    source = tmp_path / "protected.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("secret")
    with source.open("wb") as handle:
        writer.write(handle)

    with pytest.raises(DocumentReadError, match="已加密"):
        read_document_text(source)


def test_ingester_exposes_pdf_read_error_as_safe_validation_error(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with source.open("wb") as handle:
        writer.write(handle)
    ingester = DenseDocumentIngester(RagConfig(backend="qdrant"), object(), object(), object())
    def failed_parse(*args, **kwargs):
        raise ValueError("PDF parser failed")
    monkeypatch.setattr("rag_mcp.pdf_assets.parse_pdf", failed_parse)
    with pytest.raises(KnowledgeValidationError, match="PDF parser failed"):
        ingester.ingest(source, source_root=tmp_path)


def test_discovery_includes_pdf_and_word(tmp_path: Path) -> None:
    expected = [tmp_path / name for name in ("a.md", "b.pdf", "c.docx", "d.txt", "legacy.doc")]
    for path in expected:
        path.write_bytes(b"fixture")

    assert discover_documents(tmp_path) == expected


def test_malformed_docx_has_safe_error(tmp_path: Path) -> None:
    source = tmp_path / "broken.docx"
    source.write_bytes(b"not a zip archive")

    with pytest.raises(DocumentReadError, match="损坏或格式无效"):
        read_document_text(source)


def test_docx_uncompressed_size_is_bounded(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "large.docx"
    document = Document()
    document.add_paragraph("bounded")
    document.save(source)
    monkeypatch.setattr(document_readers, "MAX_DOCX_UNCOMPRESSED_BYTES", 1)

    with pytest.raises(DocumentReadError, match="解压后"):
        read_document_text(source)


def _write_text_pdf(path: Path, text: str) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_reference = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_reference})}
    )
    content = DecodedStreamObject()
    content.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(content)
    with path.open("wb") as handle:
        writer.write(handle)
