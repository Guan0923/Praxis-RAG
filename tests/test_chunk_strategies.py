"""Real document parsing and BGE-M3 tokenizer tests (dense extra required)."""

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter

from rag_mcp.chunk_strategies import ChunkingConfig, TokenBudget, get_tokenizer, sentence_ranges
from rag_mcp.chunking import chunk_document
from rag_mcp.document_readers import read_document
from test_document_readers import _write_text_pdf


@pytest.fixture(scope="module")
def tokenizer():
    pytest.importorskip("transformers")
    return get_tokenizer("BAAI/bge-m3")


def check_chunks(path, chunks, tokenizer, size):
    original = read_document(path).text
    covered = set()
    previous_end = 0
    for chunk in chunks:
        start, end = chunk.metadata["start"], chunk.metadata["end"]
        assert end > previous_end
        assert chunk.text == original[start:end]
        count = len(tokenizer.encode(chunk.text, add_special_tokens=False))
        assert 0 < count <= size
        assert chunk.metadata["token_count"] == count
        covered.update(range(start, end))
        previous_end = end
    assert all(index in covered or char.isspace() for index, char in enumerate(original))


@pytest.mark.parametrize("strategy", ["token", "sentence", "recursive", "structure"])
@pytest.mark.parametrize("text", [
    "星河设备工作正常。温度超过上限时应先断电！请检查传感器连接？\n\n" * 45,
    "Dr. Smith checked version 3.14. It works! Another sentence?\n\n" * 40,
    "Mixed中文 English café 👩🏽‍💻。下一句话有内容！\n\n" * 40,
    "无标点非常长的连续文字" * 200,
])
def test_real_tokenizer_limits_and_source_coverage(tmp_path, tokenizer, strategy, text):
    path = tmp_path / "sample.txt"
    path.write_text(text, encoding="utf-8")
    chunks = chunk_document(path, config=ChunkingConfig(strategy, 64, 12))
    assert len(chunks) > 1
    check_chunks(path, chunks, tokenizer, 64)


def test_markdown_headings_and_code_fences(tmp_path, tokenizer):
    path = tmp_path / "guide.md"
    path.write_text("# 手册\n\n开头。\n\n## 维护\n\n```python\n# not a heading\nprint('ok')\n```\n\n清洁。\n\n第二节\n------\n\n结束。", encoding="utf-8")
    document = read_document(path)
    assert [span.block.headings for span in document.spans if span.block.kind == "heading"] == [("手册",), ("手册", "维护"), ("手册", "第二节")]
    chunks = chunk_document(path, config=ChunkingConfig("structure", 512, 64))
    assert len(chunks) == 1
    assert "# not a heading" in chunks[0].text
    assert ["手册", "维护"] in chunks[0].metadata["heading_paths"]
    check_chunks(path, chunks, tokenizer, 512)


def test_pdf_merges_across_real_page_numbers(tmp_path, tokenizer):
    page = tmp_path / "page.pdf"
    _write_text_pdf(page, "First page paragraph.")
    writer = PdfWriter()
    writer.add_page(PdfReader(page).pages[0])
    writer.add_blank_page(width=612, height=792)
    _write_text_pdf(page, "Third page paragraph.")
    writer.add_page(PdfReader(page).pages[0])
    path = tmp_path / "pages.pdf"
    writer.write(path)
    chunks = chunk_document(path, config=ChunkingConfig("structure", 512, 64))
    assert len(chunks) == 1
    assert chunks[0].metadata["pages"] == [1, 3]
    check_chunks(path, chunks, tokenizer, 512)


def test_word_sections_tables_and_long_rows(tmp_path, tokenizer):
    document = Document()
    document.add_heading("设备手册", level=1)
    document.add_paragraph("正常使用。")
    document.add_heading("故障处理", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "代码", "处理方法"
    table.cell(1, 0).text, table.cell(1, 1).text = "E03", "检查传感器然后重启。" * 100
    document.add_paragraph("处理结束。")
    path = tmp_path / "guide.docx"
    document.save(path)
    chunks = chunk_document(path, config=ChunkingConfig("structure", 64, 12))
    check_chunks(path, chunks, tokenizer, 64)
    assert any(location["kind"] == "table_row" for chunk in chunks for location in chunk.metadata["locations"])
    assert any(["设备手册", "故障处理"] in chunk.metadata["heading_paths"] for chunk in chunks)


def test_sentence_boundaries_and_complete_overlap(tmp_path, tokenizer):
    text = "Dr. Smith measured 3.14. Everything works! 中文句子结束。”下一句？"
    sentences = [text[a:b].strip() for a, b in sentence_ranges(text, 0, len(text))]
    assert sentences == ["Dr. Smith measured 3.14.", "Everything works!", "中文句子结束。”", "下一句？"]
    path = tmp_path / "sentences.txt"
    path.write_text("设备正常。请检查连接。随后重启。" * 30, encoding="utf-8")
    config = ChunkingConfig("sentence", 32, 12)
    chunks = chunk_document(path, config=config)
    original = read_document(path).text
    overlaps = [original[b.metadata["start"]:a.metadata["end"]] for a, b in zip(chunks, chunks[1:]) if b.metadata["start"] < a.metadata["end"]]
    assert overlaps
    assert all(item.endswith("。") and len(tokenizer.encode(item, add_special_tokens=False)) <= 12 for item in overlaps)


def test_parameters_and_model_limit(tokenizer):
    for args in [("missing", 512, 64), ("token", 0, 0), ("token", 16, 16), ("token", 16, -1)]:
        with pytest.raises(ValueError):
            ChunkingConfig(*args)
    with pytest.raises(ValueError, match="模型"):
        TokenBudget(tokenizer, ChunkingConfig("recursive", 8192, 64))


def test_tokenizer_failure_never_falls_back(tmp_path, monkeypatch):
    path = tmp_path / "fail.txt"
    path.write_text("valid document", encoding="utf-8")
    def fail(_):
        raise ValueError("Tokenizer 加载失败")
    monkeypatch.setattr("rag_mcp.chunking.get_tokenizer", fail)
    with pytest.raises(ValueError, match="Tokenizer"):
        chunk_document(path, config=ChunkingConfig())
