"""Browser acceptance against the isolated, real Word preview service."""

import io
import os
from pathlib import Path

import httpx
import pytest
from docx import Document
from PIL import Image
from playwright.sync_api import expect, sync_playwright

from test_word_fidelity import jpeg


@pytest.mark.skipif(not os.getenv("RAG_WORD_PREVIEW_URL"), reason="Set RAG_WORD_PREVIEW_URL for real Word browser acceptance")
def test_word_evidence_desktop_mobile_and_download():
    url = os.environ["RAG_WORD_PREVIEW_URL"]
    output = Path(".test-tmp/word-browser")
    output.mkdir(parents=True, exist_ok=True)
    raw_tiff = io.BytesIO()
    Image.new("RGB", (64, 32), "green").save(raw_tiff, format="TIFF")
    doc = Document()
    doc.add_paragraph("Figure 1: ZXTIF211 original TIFF", style="Caption")
    doc.add_picture(io.BytesIO(raw_tiff.getvalue()))
    source = io.BytesIO()
    doc.save(source)
    with httpx.Client(base_url=url, headers={"X-RAG-Namespace": "manuals"}, timeout=120, trust_env=False) as client:
        response = client.post("/v1/documents", files={"files": ("tiff.docx", source.getvalue())})
        response.raise_for_status()
        assert not response.json()["failed"], response.json()
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        errors = []
        try:
            for label, width, height in (("desktop", 1440, 1000), ("mobile", 390, 844)):
                page = browser.new_page(viewport={"width": width, "height": height})
                page.on("pageerror", lambda error: errors.append(str(error)))
                page.goto(url + "/console#/workspaces/manuals/search")
                page.get_by_label("重排方式", exact=True).click()
                page.get_by_role("option", name="不重排", exact=True).click()
                page.get_by_label("检索方式", exact=True).click()
                page.get_by_role("option", name="BM25", exact=True).click()
                def search(text):
                    page.get_by_label("查询内容", exact=True).fill(text)
                    page.get_by_role("button", name="运行检索", exact=True).click()
                    expect(page.get_by_role("button", name="运行检索", exact=True)).to_be_enabled(timeout=30000)
                search("ZXIMAGE731")
                image = page.locator(".result-asset .ant-image-img").first
                expect(image).to_be_visible(timeout=30000)
                image.scroll_into_view_if_needed()
                expect(image).to_have_js_property("complete", True)
                assert image.evaluate("image => image.naturalWidth > 0")
                page.screenshot(path=str(output / f"{label}-image.png"), full_page=True)
                with page.expect_download() as download:
                    page.get_by_role("link", name="下载原图").first.click()
                saved = output / f"{label}-original.jpg"
                download.value.save_as(saved)
                assert saved.read_bytes() == jpeg()
                search("ZXTABLE942")
                expect(page.locator(".asset-table").first).to_be_visible(timeout=30000)
                expect(page.locator('.asset-table td[colspan="2"]').first).to_have_text("Merged")
                expect(page.get_by_text("ZXNESTED638", exact=True).first).to_be_visible()
                page.locator(".result-asset").first.scroll_into_view_if_needed()
                page.screenshot(path=str(output / f"{label}-table.png"), full_page=True)
                assert page.evaluate("document.documentElement.scrollWidth <= innerWidth + 1")
                search("ZXTIF211")
                expect(page.get_by_text("此图片格式不支持直接预览", exact=True)).to_be_visible(timeout=30000)
                with page.expect_download() as download:
                    page.get_by_role("link", name="下载原图").first.click()
                saved = output / f"{label}-original.tiff"
                download.value.save_as(saved)
                assert saved.read_bytes() == raw_tiff.getvalue()
                page.goto(url + "/console#/workspaces/manuals/documents/ambiguous.docx")
                expect(page.get_by_text("已跳过 2 个无法关联的对象", exact=True)).to_be_visible(timeout=30000)
                expect(page.get_by_text("Figure 7: ambiguous", exact=True).first).to_be_visible()
                page.screenshot(path=str(output / f"{label}-diagnostics.png"), full_page=True)
                assert page.evaluate("document.documentElement.scrollWidth <= innerWidth + 1")
                page.close()
        finally:
            browser.close()
        assert not errors, errors
