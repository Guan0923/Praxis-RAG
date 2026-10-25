"""Real TCP protocol tests; model inference is replaced by a controllable worker."""

from __future__ import annotations

import io
import json
import socket
import threading
import time
from contextlib import asynccontextmanager, contextmanager
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4

import anyio
import httpx2 as httpx
import pytest
import uvicorn
from mcp.client import Client
from mcp.client.streamable_http import streamable_http_client
from mcp.client.subscriptions import SubscriptionLost
from mcp.server.subscriptions import (
    PromptsListChanged,
    ResourcesListChanged,
    ResourceUpdated,
    ToolsListChanged,
)
from mcp.shared.exceptions import MCPError
from redis_stub import MemoryRedis

from rag_mcp.chunk_strategies import ChunkingConfig
from rag_mcp.config import RagConfig
from rag_mcp.document_readers import read_document_text
from rag_mcp.http_service import RagHttpSettings, create_service_app
from rag_mcp.knowledge_bases import KnowledgeBaseStore
from rag_mcp.model_api import ModelConnection
from rag_mcp.resources import document_uri
from rag_mcp.search_worker import SearchWorker
from rag_mcp.service_knowledge import (
    IndexedDocument,
    KnowledgeDocumentStore,
    RagServiceRuntime,
)
from rag_mcp.state.cache import RedisStateCache, StateCacheSettings
from rag_mcp.state.document_tasks import document_envelopes
from rag_mcp.state.settings import StateSettings
from rag_mcp.state.store import StateStores


class MemorySearchTransport:
    def __init__(self):
        self.queued = []
        self.lock = threading.Lock()

    def send(self, payload):
        with self.lock:
            self.queued.append(json.loads(json.dumps(payload)))

    def close(self):
        pass


class TestSearchConsumer:
    """Consume the test transport independently of request admission."""

    __test__ = False

    def __init__(self, runtime):
        self.runtime = runtime
        self.transport = MemorySearchTransport()
        runtime.state.search_dispatcher.transport = self.transport
        self.stop_requested = threading.Event()
        self.cursor = 0
        self.threads = []

    def start(self):
        for _ in range(2):
            thread = threading.Thread(target=self._run, daemon=True)
            self.threads.append(thread)
            thread.start()

    def _run(self):
        worker = SearchWorker(self.runtime, executor=self.runtime.search)
        while not self.stop_requested.wait(0.01):
            with self.transport.lock:
                if self.cursor >= len(self.transport.queued):
                    continue
                payload = self.transport.queued[self.cursor]
                self.cursor += 1
            worker.process(payload)

    def stop(self):
        self.stop_requested.set()
        for thread in self.threads:
            thread.join(12)
            if thread.is_alive():
                raise RuntimeError("Test search consumer failed to stop")

AUTH = {"Authorization": "Bearer test-token", "X-RAG-Namespace": "manuals"}
ADMIN = {"X-RAG-Namespace": "manuals"}


class TestIngester:
    __test__ = False

    def __init__(self, root):
        self.root = root

    def ingest(self, path, *, source_name, chunking=None):
        text = read_document_text(path)
        artifacts_dir = None
        artifacts_bucket = None
        if path.suffix in {".pdf", ".doc", ".docx"}:            # Protocol fixture supplies a saved parse; GPU parsing is tested separately.
            from rag_mcp.chunking import _stable_id
            from rag_mcp.pdf_content import PDF_PARSER_VERSION
            from rag_mcp.word_reader import WORD_PARSER_VERSION
            document_id = _stable_id("doc", source_name.casefold())
            is_word = path.suffix != ".pdf"
            artifacts_dir = self.root / f"test-assets-{uuid4().hex}"
            artifacts_dir.mkdir(parents=True)
            artifacts_bucket = "word-assets" if is_word else "pdf-assets"
            (artifacts_dir / "manifest.json").write_text(
                json.dumps({"version": WORD_PARSER_VERSION if is_word else PDF_PARSER_VERSION}), encoding="utf-8"
            )
            (artifacts_dir / "document.md").write_text(text, encoding="utf-8")
        else:
            document_id = source_name
        if path.name == "fail.txt":
            raise ValueError("controlled indexing failure")
        return IndexedDocument(document_id, 1, "test", artifacts_dir=artifacts_dir,
                               artifacts_bucket=artifacts_bucket)

    def delete(self, document_id):
        pass


class MemoryDocumentTransport:
    def __init__(self):
        self.queued = []

    def send(self, task):
        self.queued.extend(document_envelopes(task))

    def close(self):
        pass


class ProtocolRuntime(RagServiceRuntime):
    """Real document storage/reading; no model or Qdrant required for wire tests."""

    def __init__(self, root):
        self.config = RagConfig()
        self._connections = {"embedding": ModelConnection("http://unused.invalid/v1", self.config.embedding_model, vector_size=self.config.vector_size)}
        self.data_root = root
        self.state = StateStores.create(root, settings=StateSettings(
            url="sqlite+pysqlite:///" + str(root / "state.sqlite3")))
        self.state.cache_sync.backend = RedisStateCache(
            StateCacheSettings(prefix="protocol-test", enabled=True), client=MemoryRedis())
        self.state.document_dispatcher.transport = MemoryDocumentTransport()
        self.search_consumer = TestSearchConsumer(self)
        self.search_calls = []
        self._tenants = {}
        self._tenant_lock = threading.RLock()
        self._inference_lock = threading.RLock()
        self.started = threading.Event()
        self.release = threading.Event()
        self.finished = threading.Event()
        self.wire_responses = []

    def create_knowledge_base(self, namespace, chunking, embedding_model, reranker_model=""):
        settings = {"namespace": namespace, "vector_size": self.config.vector_size,
                    "embedding_model": embedding_model, "reranker_model": reranker_model,
                    "collection": self.config.collection, "tokenizer": self.config.chunk_tokenizer_model,
                    **chunking.metadata(self.config.chunk_tokenizer_model)}
        KnowledgeBaseStore(self.state).create(settings)
        return settings

    def _tenant(self, namespace):
        with self._tenant_lock:
            if namespace not in self._tenants:
                self._tenants[namespace] = SimpleNamespace(
                    state=self.state,
                    documents=KnowledgeDocumentStore(self.state, namespace),
                    ingester=TestIngester(self.data_root),
                    lock=threading.RLock(),
                )
            return self._tenants[namespace]

    def status(self):
        return {"status": "ok"}

    def search(self, namespace, arguments):
        self.search_calls.append((namespace, arguments))
        if arguments["query"] == "fail":
            raise ValueError("Controlled search failure")
        if arguments["query"] == "slow":
            self.started.set()
            self.release.wait(10)
            self.finished.set()
        return {"namespace": namespace, "query": arguments["query"], "results": []}


@contextmanager
def running_service(root: Path, *, runtime=None):
    runtime = runtime or ProtocolRuntime(root)
    if isinstance(runtime, ProtocolRuntime):
        for namespace in ("manuals", "other"):
            if KnowledgeBaseStore(runtime.state).get(namespace) is None:
                runtime.create_knowledge_base(namespace, ChunkingConfig(), runtime.config.embedding_model)
    settings = RagHttpSettings(
        data_root=root,
        token="test-token",
        mcp_allowed_hosts=("127.0.0.1:*",),
        mcp_allowed_origins=("http://127.0.0.1:*",),
    )
    app = create_service_app(runtime, settings=settings)

    class WireRecorder:
        def __init__(self, app):
            self.app = app

        async def __call__(self, scope, receive, send):
            chunks = []
            if isinstance(runtime, ProtocolRuntime) and scope["type"] == "http" and scope["path"] == "/mcp":
                runtime.wire_responses.append(chunks)

            async def record(message):
                if message["type"] == "http.response.body":
                    chunks.append(message.get("body", b""))
                await send(message)

            await self.app(scope, receive, record)

    app.add_middleware(WireRecorder)
    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    sock.listen()
    server = uvicorn.Server(uvicorn.Config(app, log_level="error", lifespan="on", timeout_graceful_shutdown=1))
    app.state.test_server = server
    thread = threading.Thread(target=server.run, kwargs={"sockets": [sock]}, daemon=True)
    thread.start()
    consumer = getattr(runtime, "search_consumer", None)
    if consumer is not None:
        consumer.start()
    deadline = time.monotonic() + 5
    while not server.started and thread.is_alive() and time.monotonic() < deadline:
        time.sleep(0.01)
    if not server.started:
        raise RuntimeError("Test service failed to start")
    try:
        yield f"http://127.0.0.1:{sock.getsockname()[1]}", app, runtime
    finally:
        if isinstance(runtime, ProtocolRuntime):
            runtime.release.set()
        server.should_exit = True
        thread.join(6)
        sock.close()
        if consumer is not None:
            consumer.stop()
        if isinstance(runtime, ProtocolRuntime):
            runtime.state.close()
        if thread.is_alive():
            raise RuntimeError("Test service failed to stop")


@asynccontextmanager
async def mcp_client(url, mode="auto", headers=None, hooks=None, message_handler=None, timeout=5):
    async with httpx.AsyncClient(headers=headers or AUTH, timeout=timeout, event_hooks=hooks or {}) as http:
        async with Client(
            streamable_http_client(url + "/mcp", http_client=http),
            mode=mode,
            read_timeout_seconds=timeout,
            message_handler=message_handler,
        ) as client:
            yield client


@pytest.mark.parametrize("mode", ["legacy", "auto", "2026-07-28"])
def test_both_protocol_lifecycles(tmp_path, mode):
    with running_service(tmp_path) as (url, app, runtime):

        async def scenario():
            requests = []
            session_ids = []

            async def record_request(request):
                if request.method == "POST":
                    requests.append(json.loads(request.content)["method"])

            async def record_response(response):
                if response.headers.get("mcp-session-id"):
                    session_ids.append(response.headers["mcp-session-id"])

            async with mcp_client(
                url, mode, hooks={"request": [record_request], "response": [record_response]}
            ) as client:
                assert client.protocol_version == ("2025-11-25" if mode == "legacy" else "2026-07-28")
                if mode != "2026-07-28":
                    assert client.server_info.name == "Praxis-RAG"
                    assert client.server_info.version == "0.4.0"
                    assert client.server_capabilities.tools.list_changed
                assert [tool.name for tool in (await client.list_tools()).tools] == [
                    "search", "get_search_result", "add_document", "update_document", "delete_document", "get_document_status"
                ]
                search_tool = next(tool for tool in (await client.list_tools()).tools if tool.name == "search")
                assert search_tool.input_schema["properties"]["wait_seconds"]["default"] == 10
                for query in ("first", "second"):
                    result = await client.call_tool("search", {"query": query})
                    assert json.loads(result.content[0].text)["query"] == query
                assert (await client.call_tool("search", {"query": ""})).is_error
                assert (await client.call_tool("missing", {})).is_error
                assert (await client.list_prompts()).prompts == []
                assert (await client.list_resources()).resources == []
            if mode == "legacy":
                assert requests[:2] == ["initialize", "notifications/initialized"]
                async with httpx.AsyncClient(headers=AUTH) as http:
                    closed = await http.post(
                        url + "/mcp",
                        headers={
                            "Mcp-Session-Id": session_ids[0],
                            "MCP-Protocol-Version": "2025-11-25",
                            "Accept": "application/json, text/event-stream",
                        },
                        json={"jsonrpc": "2.0", "id": 99, "method": "tools/list"},
                    )
                    assert closed.status_code == 404
                async with mcp_client(url, mode, hooks={"response": [record_response]}) as client:
                    await client.list_tools()
                assert session_ids[-1] != session_ids[0]
            else:
                assert "initialize" not in requests
                assert ("server/discover" in requests) == (mode == "auto")

        anyio.run(scenario)


def test_modern_subscriptions_and_real_document_pages(tmp_path):
    with running_service(tmp_path) as (url, app, runtime):

        async def scenario():
            async with httpx.AsyncClient(headers=AUTH) as http, mcp_client(url) as client:
                async with client.listen(
                    tools_list_changed=True, prompts_list_changed=True, resources_list_changed=True
                ) as sub:
                    changed = await http.patch(url + "/v1/admin/tools/search", headers=ADMIN, json={"enabled": False})
                    assert changed.json()["changed"]
                    with anyio.fail_after(2):
                        assert isinstance(await anext(sub), ToolsListChanged)
                    assert "search" not in {tool.name for tool in (await client.list_tools(cache_mode="refresh")).tools}
                    assert (await client.call_tool("search", {"query": "x"})).is_error
                    await http.patch(url + "/v1/admin/tools/search", headers=ADMIN, json={"enabled": True})
                    with anyio.fail_after(2):
                        assert isinstance(await anext(sub), ToolsListChanged)
                    template = {
                        "description": "Summarize",
                        "arguments": [{"name": "topic", "required": True}],
                        "messages": [{"role": "user", "text": "Summarize {{topic}}"}],
                    }
                    assert (
                        await http.put(url + "/v1/admin/prompts/summary", headers=ADMIN, json=template)
                    ).status_code == 200
                    with anyio.fail_after(2):
                        assert isinstance(await anext(sub), PromptsListChanged)
                    prompt = await client.get_prompt("summary", {"topic": "{{do_not_execute}}"})
                    assert prompt.messages[0].content.text == "Summarize {{do_not_execute}}"
                    with pytest.raises(MCPError):
                        await client.get_prompt("summary", {})
                    text = "段落一\n\n" + "长" * 17000 + "\n\n最后一段"
                    uploaded = await http.post(url + "/v1/documents", files={"files": ("手册.txt", text.encode())})
                    assert not uploaded.json()["failed"]
                    with anyio.fail_after(2):
                        assert isinstance(await anext(sub), ResourcesListChanged)
                    uri = (await client.list_resources(cache_mode="refresh")).resources[0].uri
                    next_uri, pieces = uri, []
                    while next_uri:
                        payload = json.loads((await client.read_resource(next_uri)).contents[0].text)
                        assert len(payload["text"]) <= 8000
                        pieces.append(payload["text"])
                        next_uri = payload["next_uri"]
                    assert "".join(pieces) == text
                    first_page = json.loads((await client.read_resource(uri)).contents[0].text)
                    async with client.listen(resource_subscriptions=[uri]) as updates:
                        await http.post(url + "/v1/documents", files={"files": ("手册.txt", b"replacement")})
                        with anyio.fail_after(2):
                            assert isinstance(await anext(sub), ResourcesListChanged)
                            assert (await anext(updates)).uri == uri
                        with pytest.raises(MCPError, match="changed"):
                            await client.read_resource(first_page["next_uri"], cache_mode="refresh")
                        await http.delete(url + "/v1/documents/手册.txt")
                        with anyio.fail_after(2):
                            assert isinstance(await anext(sub), ResourcesListChanged)
                            assert isinstance(await anext(updates), ResourceUpdated)
                        with pytest.raises(MCPError):
                            await client.read_resource(uri, cache_mode="refresh")
                    assert (await http.delete(url + "/v1/admin/prompts/summary", headers=ADMIN)).status_code == 200
                    with anyio.fail_after(2):
                        assert isinstance(await anext(sub), PromptsListChanged)

        anyio.run(scenario)


@pytest.mark.parametrize("mode", ["legacy", "2026-07-28"])
def test_cancellation_keeps_other_calls_usable(tmp_path, mode):
    with running_service(tmp_path) as (url, app, runtime):

        async def scenario():
            async with mcp_client(url, mode) as client:
                successes = []
                scope_ready = anyio.Event()
                finished = anyio.Event()
                scopes = []

                async def slow():
                    with anyio.CancelScope() as scope:
                        scopes.append(scope)
                        scope_ready.set()
                        successes.append(await client.call_tool("search", {"query": "slow"}))
                    finished.set()

                async with anyio.create_task_group() as tasks:
                    tasks.start_soon(slow)
                    await scope_ready.wait()
                    assert await anyio.to_thread.run_sync(runtime.started.wait, 2)
                    scopes[0].cancel()
                    with anyio.fail_after(2):
                        await finished.wait()
                    result = await client.call_tool("search", {"query": "still usable", "wait_seconds": 0})
                    assert not result.is_error
                    assert set(json.loads(result.content[0].text)) == {"result_id"}
                    runtime.release.set()
                    assert await anyio.to_thread.run_sync(runtime.finished.wait, 2)
                    assert not (await client.call_tool("search", {"query": "after worker"})).is_error
                    slow_id = runtime.search_consumer.transport.queued[0]["result_id"]
                    with anyio.fail_after(2):
                        while True:
                            saved = await client.call_tool("get_search_result", {"result_id": slow_id})
                            assert not saved.is_error
                            payload = json.loads(saved.content[0].text)
                            if payload.get("status") != "queued":
                                break
                            await anyio.sleep(0.02)
                    assert payload["query"] == "slow"
                    assert [args["query"] for _, args in runtime.search_calls].count("slow") == 1
                    assert successes == []

        anyio.run(scenario)


@pytest.mark.parametrize("mode", ["legacy", "2026-07-28"])
def test_queued_search_polling_is_tenant_bound_and_never_repeats_work(tmp_path, mode):
    from uuid import UUID

    with running_service(tmp_path) as (url, _, runtime):
        async def scenario():
            async with mcp_client(url, mode) as client:
                result = await client.call_tool("search", {"query": "slow", "wait_seconds": 0})
                assert not result.is_error
                admitted = json.loads(result.content[0].text)
                assert set(admitted) == {"result_id"}
                result_id = admitted["result_id"]
                assert str(UUID(result_id)) == result_id
                assert await anyio.to_thread.run_sync(runtime.started.wait, 2)
                pending = await client.call_tool("get_search_result", admitted)
                assert not pending.is_error
                assert json.loads(pending.content[0].text) == {
                    "result_id": result_id, "status": "queued", "stage": "processing",
                }
                other_headers = {**AUTH, "X-RAG-Namespace": "other"}
                async with mcp_client(url, mode, headers=other_headers) as other:
                    assert (await other.call_tool("get_search_result", admitted)).is_error
                    runtime.release.set()
                    with anyio.fail_after(2):
                        while True:
                            completed = await client.call_tool("get_search_result", admitted)
                            assert not completed.is_error
                            snapshot = json.loads(completed.content[0].text)
                            if snapshot.get("status") != "queued":
                                break
                            await anyio.sleep(0.02)
                    assert snapshot["result_id"] == result_id and snapshot["query"] == "slow"
                    assert (await other.call_tool("get_search_result", admitted)).is_error
                replay = await client.call_tool("get_search_result", admitted)
                assert json.loads(replay.content[0].text) == snapshot
                assert len(runtime.search_calls) == 1
                assert runtime.search_calls[0][0] == "manuals"
                envelopes = runtime.search_consumer.transport.queued
                assert [item["result_id"] for item in envelopes] == [result_id]
                assert envelopes[0]["schema"] == 2 and envelopes[0]["namespace"] == "manuals"
                assert envelopes[0]["arguments"]["query"] == "slow"
                assert envelopes[0]["created_at"] < envelopes[0]["deadline"] <= envelopes[0]["expires_at"]
        anyio.run(scenario)


@pytest.mark.parametrize("mode", ["legacy", "2026-07-28"])
def test_wait_validation_over_mcp_does_not_admit_tasks(tmp_path, mode):
    with running_service(tmp_path) as (url, _, runtime):
        async def scenario():
            async with mcp_client(url, mode) as client:
                for wait in (-1, 30.1, True, False, "5", None):
                    result = await client.call_tool("search", {"query": "hello", "wait_seconds": wait})
                    assert result.is_error
                assert (await client.call_tool("search", {
                    "query": "hello", "wait_seconds": 0, "namespace": "other",
                })).is_error
                assert runtime.search_calls == []
                assert runtime.search_consumer.transport.queued == []
        anyio.run(scenario)


@pytest.mark.parametrize("mode", ["legacy", "2026-07-28"])
def test_failed_search_is_an_mcp_tool_error_without_resubmission(tmp_path, mode):
    with running_service(tmp_path) as (url, _, runtime):
        async def scenario():
            async with mcp_client(url, mode) as client:
                failed = await client.call_tool("search", {"query": "fail"})
                assert failed.is_error
                assert len(runtime.search_calls) == 1
                result_id = runtime.search_consumer.transport.queued[0]["result_id"]
                assert (await client.call_tool("get_search_result", {"result_id": result_id})).is_error
                assert len(runtime.search_calls) == 1
                assert [item["result_id"] for item in runtime.search_consumer.transport.queued] == [result_id]
        anyio.run(scenario)


def test_search_deadline_is_an_mcp_tool_error_without_inference(tmp_path):
    with running_service(tmp_path) as (url, _, runtime):
        runtime.search_consumer.stop()

        async def scenario():
            async with mcp_client(url) as client:
                result = await client.call_tool("search", {"query": "never run", "wait_seconds": 0})
                assert not result.is_error
                admitted = json.loads(result.content[0].text)
                assert set(admitted) == {"result_id"}
                pending = await client.call_tool("get_search_result", admitted)
                assert not pending.is_error
                assert json.loads(pending.content[0].text) == {
                    **admitted, "status": "queued", "stage": "queued",
                }
                runtime.state.search_tasks.transient.update(
                    "search-result", admitted["result_id"], lambda row: {**row, "deadline": 0})
                assert (await client.call_tool("get_search_result", admitted)).is_error
                assert runtime.search_calls == []
        anyio.run(scenario)


def test_wait_timeout_leaves_background_search_running(tmp_path):
    with running_service(tmp_path) as (url, _, runtime):
        async def scenario():
            async with mcp_client(url) as client:
                result = await client.call_tool("search", {"query": "slow", "wait_seconds": 0.05})
                assert not result.is_error
                admitted = json.loads(result.content[0].text)
                assert set(admitted) == {"result_id"}
                assert await anyio.to_thread.run_sync(runtime.started.wait, 2)
                assert not runtime.finished.is_set()
                runtime.release.set()
                with anyio.fail_after(2):
                    while True:
                        polled = await client.call_tool("get_search_result", admitted)
                        assert not polled.is_error
                        saved = json.loads(polled.content[0].text)
                        if saved.get("status") != "queued":
                            break
                        await anyio.sleep(0.02)
                assert saved["query"] == "slow" and saved["result_id"] == admitted["result_id"]
                assert len(runtime.search_calls) == 1
                assert [item["result_id"] for item in runtime.search_consumer.transport.queued] == [admitted["result_id"]]
        anyio.run(scenario)


def test_http_security_and_local_management(tmp_path):
    with running_service(tmp_path) as (url, app, runtime):

        async def scenario():
            request = {
                "jsonrpc": "2.0",
                "id": 1,
                "method": "initialize",
                "params": {
                    "protocolVersion": "2025-11-25",
                    "capabilities": {},
                    "clientInfo": {"name": "test", "version": "1"},
                },
            }
            base = {**AUTH, "Accept": "application/json, text/event-stream"}
            async with httpx.AsyncClient() as http:
                for headers, status in [
                    ({**base, "Authorization": "Bearer wrong"}, 401),
                    ({**base, "X-RAG-Namespace": "../escape"}, 400),
                    ({**base, "Host": "evil.example"}, 421),
                    ({**base, "Origin": "https://evil.example"}, 403),
                    ({**base, "Authorization": "Bearer admin-token"}, 401),
                ]:
                    assert (await http.post(url + "/mcp", headers=headers, json=request)).status_code == status
                assert (await http.get(url + "/v1/admin/tools", headers=AUTH)).status_code == 200
                assert (await http.get(url + "/v1/admin/tools", headers=ADMIN)).status_code == 200
                assert not runtime._tenants  # configuration does not initialize retrieval dependencies
                assert (
                    await http.patch(url + "/v1/admin/tools/nope", headers=ADMIN, json={"enabled": True})
                ).status_code == 404
                assert (
                    await http.patch(url + "/v1/admin/tools/search", headers=ADMIN, json={"code": "unsafe"})
                ).status_code == 422

        anyio.run(scenario)


def test_local_management_needs_no_credential(tmp_path):
    with running_service(tmp_path) as (url, _, _):
        assert httpx.get(url + "/v1/admin/tools", headers=ADMIN).status_code == 200


def test_namespace_sessions_and_subscription_filters(tmp_path, monkeypatch):
    with running_service(tmp_path) as (url, app, runtime):

        async def scenario():
            other = {**AUTH, "X-RAG-Namespace": "other"}
            async with httpx.AsyncClient() as http, mcp_client(url) as first, mcp_client(url, headers=other) as second:
                released = threading.Event()
                bus = app.state.mcp_endpoint.peers["manuals"].bus
                subscribe = bus.subscribe

                def watch_release(listener):
                    remove = subscribe(listener)

                    def unsubscribe():
                        remove()
                        released.set()

                    return unsubscribe

                monkeypatch.setattr(bus, "subscribe", watch_release)
                async with (
                    first.listen(tools_list_changed=True) as tools,
                    second.listen(tools_list_changed=True) as others,
                ):
                    await http.patch(url + "/v1/admin/tools/search", headers=ADMIN, json={"description": "updated"})
                    with anyio.fail_after(2):
                        assert isinstance(await anext(tools), ToolsListChanged)
                    with anyio.move_on_after(0.15) as timeout:
                        await anext(others)
                    assert timeout.cancel_called
                    prompt = {"messages": [{"role": "user", "text": "Hello"}]}
                    await http.put(url + "/v1/admin/prompts/hello", headers=ADMIN, json=prompt)
                    with anyio.move_on_after(0.15) as timeout:
                        await anext(tools)
                    assert timeout.cancel_called
                    assert not (await second.list_prompts()).prompts
                assert await anyio.to_thread.run_sync(released.wait, 2)
                with pytest.raises(MCPError):
                    await first.read_resource(document_uri("other", "secret.txt"))
                with pytest.raises(MCPError):
                    await first.read_resource("rag://knowledge/manuals/%2e%2e%2fsecret.txt")
            ids = []

            async def record(response):
                if response.headers.get("mcp-session-id"):
                    ids.append(response.headers["mcp-session-id"])

            async with mcp_client(url, "legacy", hooks={"response": [record]}) as client:
                async with httpx.AsyncClient(headers=other) as http:
                    response = await http.post(
                        url + "/mcp",
                        headers={
                            "Mcp-Session-Id": ids[0],
                            "MCP-Protocol-Version": "2025-11-25",
                            "Accept": "application/json, text/event-stream",
                        },
                        json={"jsonrpc": "2.0", "id": 50, "method": "tools/list"},
                    )
                    assert response.status_code == 404
                assert not (await client.call_tool("search", {"query": "still here"})).is_error

        anyio.run(scenario)


def test_noop_and_failed_writes_do_not_notify(tmp_path, monkeypatch):
    with running_service(tmp_path) as (url, app, runtime):

        async def scenario():
            async with httpx.AsyncClient(headers=AUTH) as http, mcp_client(url) as client:
                await http.post(url + "/v1/documents", files={"files": ("same.txt", b"same")})
                async with client.listen(
                    tools_list_changed=True, prompts_list_changed=True, resources_list_changed=True
                ) as sub:
                    assert not (
                        await http.patch(url + "/v1/admin/tools/search", headers=ADMIN, json={"enabled": True})
                    ).json()["changed"]
                    assert (
                        await http.post(url + "/v1/documents", files={"files": ("same.txt", b"same")})
                    ).status_code == 200
                    assert (await http.post(url + "/v1/documents", files={"files": ("fail.txt", b"bad")})).json()[
                        "failed"
                    ]
                    from rag_mcp.state.repositories import StateError

                    def fail_mutate(*_args, **_kwargs):
                        raise StateError("controlled database failure")

                    monkeypatch.setattr(app.state.mcp_endpoint.state.catalogs, "mutate", fail_mutate)
                    failed = await http.patch(url + "/v1/admin/tools/search", headers=ADMIN, json={"enabled": False})
                    assert failed.status_code == 503
                    assert (await client.list_tools(cache_mode="refresh")).tools
                    with anyio.move_on_after(0.2) as timeout:
                        await anext(sub)
                    assert timeout.cancel_called
                    with anyio.move_on_after(0.2) as timeout:
                        await anext(sub)
                    assert timeout.cancel_called

        anyio.run(scenario)


def test_configuration_survives_service_restart(tmp_path):
    template = {"messages": [{"role": "user", "text": "Persisted"}]}
    with running_service(tmp_path) as (url, _, _):
        assert httpx.patch(url + "/v1/admin/tools/search", headers=ADMIN, json={"enabled": False}).status_code == 200
        assert httpx.put(url + "/v1/admin/prompts/saved", headers=ADMIN, json=template).status_code == 200
    with running_service(tmp_path) as (url, _, runtime):
        tools = httpx.get(url + "/v1/admin/tools", headers=ADMIN).json()["tools"]
        prompts = httpx.get(url + "/v1/admin/prompts", headers=ADMIN).json()["prompts"]
        assert not tools[0]["enabled"]
        assert prompts[0]["messages"][0]["text"] == "Persisted"
        assert not runtime._tenants


def test_legacy_change_notifications_and_unsubscribe(tmp_path):
    with running_service(tmp_path) as (url, app, runtime):

        async def scenario():
            events_send, events = anyio.create_memory_object_stream(20)

            async def record(message):
                if hasattr(message, "method") and message.method.startswith("notifications/"):
                    await events_send.send(message)

            async with (
                httpx.AsyncClient(headers=AUTH) as http,
                mcp_client(url, "legacy", message_handler=record) as client,
            ):
                await http.post(url + "/v1/documents", files={"files": ("legacy.txt", b"first")})
                with anyio.fail_after(2):
                    assert (await events.receive()).method == "notifications/resources/list_changed"
                uri = document_uri("manuals", "legacy.txt")
                await client.subscribe_resource(uri)
                await http.post(url + "/v1/documents", files={"files": ("legacy.txt", b"second")})
                with anyio.fail_after(2):
                    methods = [(await events.receive()).method, (await events.receive()).method]
                assert set(methods) == {"notifications/resources/list_changed", "notifications/resources/updated"}
                await client.unsubscribe_resource(uri)
                await http.delete(url + "/v1/documents/legacy.txt")
                with anyio.fail_after(2):
                    assert (await events.receive()).method == "notifications/resources/list_changed"
                with anyio.move_on_after(0.15) as timeout:
                    await events.receive()
                assert timeout.cancel_called
                await http.patch(url + "/v1/admin/tools/search", headers=ADMIN, json={"description": "changed"})
                with anyio.fail_after(2):
                    assert (await events.receive()).method == "notifications/tools/list_changed"
                await http.put(
                    url + "/v1/admin/prompts/p", headers=ADMIN, json={"messages": [{"role": "user", "text": "Hi"}]}
                )
                with anyio.fail_after(2):
                    assert (await events.receive()).method == "notifications/prompts/list_changed"

        anyio.run(scenario)


def test_word_and_pdf_resources_over_http(tmp_path):
    from docx import Document
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    word = Document()
    word.add_paragraph("Word paragraph text")
    word_bytes = io.BytesIO()
    word.save(word_bytes)
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 250 Td (PDF paragraph text) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    pdf_bytes = io.BytesIO()
    writer.write(pdf_bytes)
    with running_service(tmp_path) as (url, _, _):

        async def scenario():
            async with httpx.AsyncClient(headers=AUTH) as http, mcp_client(url) as client:
                for name, blob, expected in [
                    ("read.docx", word_bytes.getvalue(), "Word paragraph text"),
                    ("read.pdf", pdf_bytes.getvalue(), "PDF paragraph text"),
                ]:
                    response = await http.post(url + "/v1/documents", files={"files": (name, blob)})
                    assert not response.json()["failed"]
                    result = await client.read_resource(document_uri("manuals", name))
                    assert expected in json.loads(result.contents[0].text)["text"]

        anyio.run(scenario)


def test_modern_protocol_headers_are_validated(tmp_path):
    with running_service(tmp_path) as (url, _, _):

        async def scenario():
            captured = []

            async def record(request):
                if request.method == "POST":
                    captured.append((dict(request.headers), json.loads(request.content)))

            async with mcp_client(url, "2026-07-28", hooks={"request": [record]}) as client:
                await client.list_tools()
            headers, body = captured[0]
            headers.pop("content-length", None)
            async with httpx.AsyncClient() as http:
                correct = await http.post(url + "/mcp", headers=headers, json=body)
                assert correct.status_code == 200
                for changes in [
                    {"mcp-protocol-version": "2099-01-01"},
                    {"mcp-method": "tools/call"},
                ]:
                    response = await http.post(url + "/mcp", headers={**headers, **changes}, json=body)
                    assert response.status_code == 400

        anyio.run(scenario)


def test_registry_dispatches_an_additional_local_tool():
    from mcp import types
    from rag_mcp.server import RegisteredTool, ToolRegistry

    calls = []

    def execute(arguments):
        calls.append(arguments["name"])
        return {"document": arguments["name"]}

    registry = ToolRegistry(
        [
            RegisteredTool(
                types.Tool(name="test_document", input_schema={"type": "object"}),
                dict,
                execute,
            )
        ]
    )
    result = registry.invoke("test_document", {"name": "manual.txt"})
    assert calls == ["manual.txt"]
    assert json.loads(result.content[0].text) == {"document": "manual.txt"}


def test_shutdown_releases_open_subscription(tmp_path):
    with running_service(tmp_path) as (url, app, _):

        async def scenario():
            async with mcp_client(url) as client:
                async with client.listen(tools_list_changed=True) as sub:
                    app.state.test_server.should_exit = True
                    with anyio.fail_after(4):
                        with pytest.raises((SubscriptionLost, StopAsyncIteration)):
                            await anext(sub)

        anyio.run(scenario)


def test_partial_upload_notifications_and_resource_listing_pages(tmp_path):
    with running_service(tmp_path) as (url, _, _):

        async def scenario():
            async with httpx.AsyncClient(headers=AUTH) as http, mcp_client(url) as client:
                async with client.listen(resources_list_changed=True) as sub:
                    result = await http.post(
                        url + "/v1/documents",
                        files=[
                            ("files", ("ok.txt", b"ok")),
                            ("files", ("fail.txt", b"bad")),
                        ],
                    )
                    assert [item["name"] for item in result.json()["imported"]] == ["ok.txt"]
                    assert [item["name"] for item in result.json()["failed"]] == ["fail.txt"]
                    with anyio.fail_after(2):
                        assert isinstance(await anext(sub), ResourcesListChanged)
                for batch in range(5):
                    result = await http.post(
                        url + "/v1/documents",
                        files=[("files", (f"{batch}-{number}.txt", f"page {batch}-{number}".encode())) for number in range(10)],
                    )
                    assert not result.json()["failed"]
                page1 = await client.list_resources(cache_mode="refresh")
                assert len(page1.resources) == 50 and page1.next_cursor
                page2 = await client.list_resources(cursor=page1.next_cursor)
                assert len(page2.resources) == 1 and page2.next_cursor is None
                assert (await client.list_resource_templates()).resource_templates
                uri = document_uri("manuals", "ok.txt")
                for bad_uri in [uri + "?page=0", uri + "?page=2", uri + "?page=999&version=old"]:
                    with pytest.raises(MCPError):
                        await client.read_resource(bad_uri)
                with pytest.raises(MCPError):
                    async with client.listen(resource_subscriptions=[uri + "?page=1"]):
                        pass

        anyio.run(scenario)


@pytest.mark.parametrize("kind", ["document", "search"])
@pytest.mark.parametrize("failure", ["broker_missing", "broker_error", "redis_disabled", "redis_error"])
def test_mcp_queue_admission_fails_closed_without_result_ids(tmp_path, kind, failure):
    with running_service(tmp_path) as (url, _, runtime):
        dispatcher = getattr(runtime.state, kind + "_dispatcher")
        transport = dispatcher.transport
        if failure == "broker_missing":
            dispatcher.transport = None
        elif failure == "broker_error":
            def reject(_payload):
                raise ConnectionError("controlled broker failure")
            transport.send = reject
        elif failure == "redis_disabled":
            runtime.state.cache_sync.backend = RedisStateCache(StateCacheSettings(enabled=False))
        else:
            runtime.state.cache_sync.backend._connection().broken = True

        async def scenario():
            async with mcp_client(url) as client:
                arguments = ({"name": "unaccepted.txt", "content": "not published"}
                             if kind == "document" else {"query": "not accepted", "wait_seconds": 0})
                result = await client.call_tool("add_document" if kind == "document" else "search", arguments)
                assert result.is_error
                assert runtime.search_calls == []
                assert transport.queued == []
                assert runtime.state.documents.list("manuals") == []
        anyio.run(scenario)


def test_mcp_redis_eviction_cannot_restore_snapshot_or_reexecute_search(tmp_path):
    with running_service(tmp_path) as (url, _, runtime):
        async def scenario():
            async with mcp_client(url) as client:
                completed = await client.call_tool("search", {"query": "frozen"})
                assert not completed.is_error
                saved = json.loads(completed.content[0].text)
                assert saved["snapshot_available"] is True
                key = runtime.state.results.transient._key("search-result", saved["result_id"])
                assert runtime.state.cache_sync.backend._connection().delete(key) == 1
                unavailable = await client.call_tool("get_search_result", {"result_id": saved["result_id"]})
                assert unavailable.is_error
                assert len(runtime.search_calls) == 1
        anyio.run(scenario)
