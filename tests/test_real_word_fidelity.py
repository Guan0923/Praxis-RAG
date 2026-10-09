"""Opt-in real LibreOffice, Docling Vision, embedding HTTP and Qdrant acceptance."""

import base64
import io
import json
import os
import subprocess
import threading
from pathlib import Path
from uuid import uuid4

import anyio
import httpx
import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont

from rag_mcp.config import RagConfig
from rag_mcp.resources import document_uri
from rag_mcp.service_knowledge import RagServiceRuntime
from rag_mcp.word_reader import WORD_PARSER_VERSION
from test_mcp_http import AUTH, mcp_client, running_service
from test_word_fidelity import jpeg, math_xml


def make_documents(root):
    captioned = Document()
    captioned.add_heading("Word acceptance", 1)
    captioned.add_paragraph("Figure 1: ZXIMAGE731 red indicator", style="Caption")
    captioned.add_picture(io.BytesIO(jpeg()))
    captioned.add_paragraph("Image section ends here.")
    captioned.add_paragraph("Table 1: ZXTABLE942 pump limits", style="Caption")
    table = captioned.add_table(rows=3, cols=3)
    for cell, text in zip(table.rows[0].cells, ("Model", "Minimum", "Maximum")):
        cell.text = text
    table.cell(1, 0).text = "ZQ-25"
    table.cell(1, 2).text = "100"
    table.cell(2, 0).merge(table.cell(2, 1)).text = "Merged"
    table.cell(2, 2).add_table(rows=1, cols=1).cell(0, 0).text = "ZXNESTED638"
    captioned.add_paragraph("Note: values in kPa")
    captioned.add_paragraph("Formula:")._p.append(math_xml('<m:f><m:num><m:r><m:t>x</m:t></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f>'))
    captioned.sections[0].header.paragraphs[0].text = "ZXHEADER554"
    captioned.save(root / "captioned.docx")
    vision = Document()
    vision.add_heading("Uncaptioned indicator", 1)
    image = Image.new("RGB", (800, 520), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=36)
    draw.rectangle((30, 30, 770, 490), fill="#e2e7eb", outline="black", width=4)
    draw.text((100, 60), "PUMP CONTROL PANEL", font=font, fill="black")
    draw.ellipse((100, 180, 300, 380), fill="#27b54a", outline="black", width=5)
    draw.text((108, 410), "RUNNING", font=font, fill="black")
    draw.rectangle((410, 200, 710, 340), fill="#202124")
    draw.text((445, 245), "120 kPa", font=font, fill="#7aff8a")
    original = io.BytesIO()
    image.save(original, format="PNG")
    vision.add_picture(io.BytesIO(original.getvalue()))
    vision.save(root / "vision.docx")
    legacy = Document()
    legacy.add_heading("Legacy manual", 1)
    legacy.add_paragraph("ZXLEGACY816 readable legacy Word")
    legacy.add_table(rows=1, cols=2).cell(0, 0).text = "Legacy table"
    legacy.save(root / "legacy.docx")
    ambiguous = Document()
    ambiguous.add_paragraph("ZXSKIP017 keep body")
    ambiguous.add_picture(io.BytesIO(jpeg()))
    ambiguous.add_paragraph("Figure 7: ambiguous", style="Caption")
    ambiguous.add_picture(io.BytesIO(jpeg()))
    ambiguous.save(root / "ambiguous.docx")


@pytest.mark.skipif(os.getenv("RAG_REAL_WORD") != "1", reason="Set RAG_REAL_WORD=1 for real Word acceptance")
def test_real_word_ingestion_search_and_mcp():
    root = Path(".test-tmp/word-real") / uuid4().hex[:8]
    root.mkdir(parents=True)
    root = root.resolve()
    make_documents(root)
    converter = Path(os.environ["RAG_WORD_CONVERTER_PATH"])
    conversion = subprocess.run([str(converter), f"-env:UserInstallation={(root / 'lo-fixture').as_uri()}",
        "--headless", "--convert-to", "doc:MS Word 97", "--outdir", str(root), str(root / "legacy.docx")],
        capture_output=True, timeout=120, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    assert conversion.returncode == 0 and (root / "legacy.doc").is_file(), conversion.stderr
    assert (root / "legacy.doc").read_bytes().startswith(bytes.fromhex("D0CF11E0A1B11AE1"))
    config = RagConfig(backend="qdrant", qdrant_path=root / "qdrant", collection="word_acceptance",
        word_converter_path=converter, pdf_parser_python=Path(os.environ["RAG_PDF_PARSER_PYTHON"]),
        pdf_models_path=Path(os.environ["RAG_PDF_MODELS_PATH"]))
    runtime = RagServiceRuntime(config, root / "data")
    runtime.release, runtime.wire_responses = threading.Event(), []
    report = {"root": str(root), "version": WORD_PARSER_VERSION, "documents": []}
    from rag_mcp.chunk_strategies import ChunkingConfig
    runtime.create_knowledge_base("manuals", ChunkingConfig(strategy="character", chunk_size=700, chunk_overlap=100),
        os.environ.get("RAG_TEST_EMBEDDING_MODEL", "bge-m3:latest"),
        embedding={"base_url": os.environ["RAG_MODEL_PROVIDER_URL"]})
    try:
        with running_service(root / "data", runtime=runtime) as (url, _, _):
            with httpx.Client(base_url=url, headers=AUTH, timeout=900, trust_env=False) as client:
                def upload(name):
                    with (root / name).open("rb") as stream:
                        response = client.post("/v1/documents", files={"files": (name, stream)}, data={"strategy": "character", "chunk_size": "700", "chunk_overlap": "100"})
                    response.raise_for_status()
                    body = response.json()
                    assert not body["failed"], body
                    return body
                for name in ("vision.docx", "captioned.docx", "legacy.doc", "ambiguous.docx"):
                    result = upload(name)
                    report["documents"].append(result["imported"][0])
                assert upload("captioned.docx")["skipped"]
                assert report["documents"][-1]["skipped_object_count"] == 2
                knowledge = root / "data/manuals/knowledge"
                captioned = report["documents"][1]
                folder = knowledge / ".word-assets" / captioned["document_id"]
                saved = (folder / "document.md").read_text(encoding="utf-8")
                assert "ZXNESTED638" in saved and "ZXHEADER554" in saved and r"\frac{x}{2}" in saved
                assert client.get("/v1/documents/captioned.docx/content").text == saved
                vision_folder = knowledge / ".word-assets" / report["documents"][0]["document_id"]
                vision = json.loads((vision_folder / "manifest.json").read_text(encoding="utf-8"))["assets"][0]
                assert vision["generated_description"] and not vision["caption"]
                assert sum(term in vision["generated_description"].lower() for term in ("pump", "control", "panel", "green", "pressure", "indicator")) >= 2, vision["generated_description"]
                report["vision_description"] = vision["generated_description"]
                def search(query, mode="bm25", source=None):
                    response = client.post("/v1/search", json={"query": query, "retrieval": mode, "rerank": "none", "candidate_k": 30, "top_k": 10,
                                          "filters": {"source": source} if source else {}})
                    response.raise_for_status()
                    return response.json()["results"]
                image = next(hit["asset"] for hit in search("ZXIMAGE731") if hit.get("asset", {}).get("type") == "image")
                assert image["mime_type"] == "image/jpeg" and base64.b64decode(image["data"]) == jpeg()
                assert "page" not in image and "bbox" not in image
                table = next(hit["asset"] for hit in search("ZXTABLE942") if hit.get("asset", {}).get("type") == "table")
                assert table["rows"][1] == ["ZQ-25", "", "100"] and table["nested_tables"]
                assert table["footnotes"] == ["Note: values in kPa"]
                assert any(hit.get("asset", {}).get("type") == "image" for hit in search(vision["generated_description"], "hybrid", "vision.docx"))
                assert any("ZXLEGACY816" in hit.get("text", "") for hit in search("ZXLEGACY816"))
                before = (knowledge / "captioned.docx").read_bytes()
                failed = client.post("/v1/documents", files={"files": ("captioned.docx", b"broken zip")}).json()
                assert failed["failed"] and (knowledge / "captioned.docx").read_bytes() == before
                assert (folder / "document.md").read_text(encoding="utf-8") == saved
                assert any(hit.get("asset", {}).get("type") == "image" for hit in search("ZXIMAGE731"))

            async def mcp_check():
                async with mcp_client(url, timeout=120) as client:
                    result = await client.call_tool("search", {"query": "ZXIMAGE731", "retrieval": "bm25", "rerank": "none", "candidate_k": 30, "top_k": 10})
                    assert not result.is_error
                    images = [part for part in result.content if part.type == "image"]
                    assert images and base64.b64decode(images[0].data) == jpeg()
                    resource = await client.read_resource(document_uri("manuals", "captioned.docx"))
                    assert "ZXHEADER554" in json.loads(resource.contents[0].text)["text"]
            anyio.run(mcp_check)
            with httpx.Client(base_url=url, headers=AUTH, timeout=120, trust_env=False) as client:
                replacement = Document()
                replacement.add_paragraph("ZXREPLACED396 new document")
                stream = io.BytesIO()
                replacement.save(stream)
                response = client.post("/v1/documents", files={"files": ("captioned.docx", stream.getvalue())}).json()
                assert response["imported"] and not response["failed"]
                assert not list((folder / "images").glob("*"))
                response = client.post("/v1/search", json={"query": "ZXIMAGE731", "retrieval": "bm25", "rerank": "none"}).json()
                assert not response["results"]
                assert client.delete("/v1/documents/captioned.docx").status_code == 200
                assert not folder.exists()
            report["checks"] = {"doc_conversion": True, "real_vision": True, "original_image_bytes": True,
                "table_structure": True, "http_and_mcp": True, "failed_replace_preserved_old": True,
                "replacement_removed_old_assets": True, "delete_removed_assets": True}
    finally:
        (root / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        runtime._client.close()
        print("WORD_REAL_REPORT", root / "report.json")
