"""Opt-in real BGE-M3 + BM25 + local Qdrant acceptance test."""

import hashlib
import io
import json
import os

import pytest
from docx import Document

from rag_mcp.chunk_strategies import ChunkingConfig
from rag_mcp.config import RagConfig
from rag_mcp.service_knowledge import RagServiceRuntime
from test_document_readers import _write_text_pdf
from model_api_test_support import configure_api_runtime


@pytest.mark.skipif(os.getenv("RAG_REAL_SMOKE") != "1", reason="Set RAG_REAL_SMOKE=1 for real model inference")
def test_real_model_ingestion_and_retrieval(tmp_path):
    config = RagConfig(backend="qdrant", qdrant_path=tmp_path / "qdrant",
                       collection="chunking_acceptance", embedding_device="cuda:0")
    runtime = RagServiceRuntime(config, tmp_path / "sources")
    source = "E03 means a temperature sensor fault. Disconnect power, check the sensor cable, then restart. "
    markdown = tmp_path / "manual.md"
    markdown.write_text("# Maintenance\n\n## E03\n\n" + source * 10, encoding="utf-8")
    plain = tmp_path / "notes.txt"
    plain.write_text("星河设备维护说明。\n\n" + source * 10, encoding="utf-8")
    pdf = tmp_path / "manual.pdf"
    _write_text_pdf(pdf, source * 4)
    word = tmp_path / "manual.docx"
    document = Document()
    document.add_heading("Maintenance", level=1)
    document.add_heading("E03", level=2)
    document.add_paragraph(source * 10)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "E03", "Check sensor cable"
    document.save(word)
    report = []
    try:
        configure_api_runtime(runtime)
        for strategy in ("character", "token", "sentence", "recursive", "structure"):
            namespace = f"smoke-{strategy}"
            chunking = ChunkingConfig(strategy, 700, 100) if strategy == "character" else ChunkingConfig(strategy, 64, 12)
            runtime.create_knowledge_base(namespace, chunking, runtime.config.embedding_model, runtime.config.reranker_model, **runtime.provider_settings)
            files = [(path.name, io.BytesIO(path.read_bytes())) for path in (markdown, plain, pdf, word)]
            result = runtime.upload(namespace, files)
            assert not result["failed"], result
            assert len(result["imported"]) == 4
            for retrieval in ("embedding", "bm25", "hybrid"):
                search = runtime.search(namespace, {"query": "E03 sensor cable restart", "retrieval": retrieval,
                                                    "rerank": "none", "candidate_k": 10, "top_k": 5})
                assert search["results"], search
                assert any("E03" in item["text"] for item in search["results"])
            fingerprint = hashlib.sha256(markdown.read_bytes()).hexdigest()
            checked = runtime.check_documents(namespace, [{"name": "renamed.md", "sha256": fingerprint}])
            assert checked["files"][0]["document"]["name"] == markdown.name
            duplicate = runtime.upload(namespace, [("renamed.md", io.BytesIO(markdown.read_bytes()))])
            assert len(duplicate["skipped"]) == 1 and not duplicate["imported"]
            report.append({"strategy": strategy, "documents": 4,
                           "chunks": sum(item["chunk_count"] for item in result["imported"]),
                           "retrieval": ["embedding", "bm25", "hybrid"], "rename_dedup": True})
        print("REAL_ACCEPTANCE=" + json.dumps(report))
    finally:
        runtime._client.close()
