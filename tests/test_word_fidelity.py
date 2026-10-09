"""Real OOXML fixtures; no replacement for the parser under test."""

import io
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from PIL import Image

from rag_mcp.word_reader import NS, parse_docx, render_word, formula


def jpeg():
    stream = io.BytesIO()
    Image.new("RGB", (48, 32), "red").save(stream, format="JPEG")
    return stream.getvalue()


def save(document, root, name="fixture.docx"):
    path = root / name
    document.save(path)
    return path


def math_xml(inner):
    return parse_xml(f'<m:oMath xmlns:m="{NS["m"]}">{inner}</m:oMath>')


def test_controls_nested_tables_empty_columns_and_merges(tmp_path):
    doc = Document()
    doc.add_heading("Manual", 1)
    table = doc.add_table(rows=3, cols=3)
    for cell, value in zip(table.rows[0].cells, ("Item", "Minimum", "Maximum")):
        cell.text = value
    table.cell(1, 0).text = "Pressure"
    table.cell(1, 2).text = "100"
    table.cell(2, 0).merge(table.cell(2, 1)).text = "Merged"
    table.cell(2, 2).add_table(rows=1, cols=1).cell(0, 0).text = "Nested value"
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    table._tbl.addprevious(control)
    content.append(table._tbl)
    control.append(content)
    result = parse_docx(save(doc, tmp_path))
    outer, nested = result["assets"]
    assert outer["table"]["rows"][1] == ["Pressure", "", "100"]
    assert outer["table"]["rows"][2] == ["Merged", "", ""]
    assert outer["table"]["cells"][-2]["col_span"] == 2
    assert nested["table"]["parent_cell"] == {"asset_id": outer["asset_id"], "row": 2, "col": 2}
    assert "Nested value" in render_word(result)


def test_vertical_merge_only_keeps_anchor(tmp_path):
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(2, 0)).text = "Only once"
    result = parse_docx(save(doc, tmp_path))
    table = result["assets"][0]["table"]
    assert table["cells"][0]["row_span"] == 3
    assert [r[0] for r in table["rows"]] == ["Only once", "", ""]


@pytest.mark.parametrize(("inner", "expected"), [
    ('<m:f><m:num><m:r><m:t>x</m:t></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f>', r'\frac{x}{2}'),
    ('<m:sSup><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>', '{x}^{2}'),
    ('<m:rad><m:e><m:r><m:t>x</m:t></m:r></m:e></m:rad>', r'\sqrt{x}'),
    ('<m:m><m:mr><m:e><m:r><m:t>1</m:t></m:r></m:e><m:e><m:r><m:t>2</m:t></m:r></m:e></m:mr></m:m>', r'\begin{matrix} 1 & 2 \end{matrix}'),
    ('<m:r><m:rPr><m:scr m:val="double-struck"/></m:rPr><m:t>R</m:t></m:r>', r'\mathbb{R}'),
])
def test_formula_structure(inner, expected):
    assert formula(math_xml(inner)) == expected


def test_unknown_and_incomplete_formula_fail():
    with pytest.raises(ValueError, match="Unsupported"):
        formula(math_xml('<m:unknown><m:r><m:t>x</m:t></m:r></m:unknown>'))
    with pytest.raises(ValueError, match="Incomplete"):
        formula(math_xml('<m:f><m:num><m:r><m:t>x</m:t></m:r></m:num></m:f>'))
    assert formula(math_xml('<m:f><m:num><m:r><m:t/></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f>')) == r'\frac{}{2}'
    with pytest.raises(ValueError, match="operator"):
        formula(math_xml('<m:nary><m:naryPr><m:chr m:val="?"/></m:naryPr><m:e><m:r><m:t>x</m:t></m:r></m:e></m:nary>'))


def test_vision_refusal_is_not_a_description():
    from rag_mcp.word_vision_worker import validate_description
    with pytest.raises(ValueError, match="Vision did not describe"):
        validate_description("I'm sorry, but I cannot provide a description of this image.")
    assert validate_description("A red indicator is lit on the control panel.")


def test_original_image_bytes_caption_note_and_inline_order(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("Before image")
    shape = paragraph.add_run().add_picture(io.BytesIO(jpeg()))
    shape._inline.docPr.set("title", "Figure 1")
    paragraph.add_run("After image")
    doc.add_paragraph("Figure 1: Red status lamp", style="Caption")
    doc.add_paragraph("Note: keep this lamp visible")
    doc.add_paragraph("Tail")
    result = parse_docx(save(doc, tmp_path), tmp_path)
    asset = result["assets"][0]
    assert (tmp_path / asset["image_file"]).read_bytes() == jpeg()
    assert asset["mime_type"] == "image/jpeg"
    assert asset["caption"] == "Figure 1: Red status lamp"
    assert asset["footnotes"][0]["text"] == "Note: keep this lamp visible"
    kinds = [(b["kind"], b["text"]) for b in result["blocks"]]
    assert kinds[:3] == [("paragraph", "Before image"), ("picture", ""), ("paragraph", "After image")]


def test_uncaptioned_image_remains_for_vision(tmp_path):
    doc = Document()
    doc.add_picture(io.BytesIO(jpeg()))
    result = parse_docx(save(doc, tmp_path), tmp_path)
    assert result["assets"][0]["caption"] == ""
    assert result["skipped_object_count"] == 0


def test_ambiguous_caption_skips_objects_and_keeps_text(tmp_path):
    doc = Document()
    doc.add_picture(io.BytesIO(jpeg()))
    doc.add_paragraph("Figure 7: ambiguous", style="Caption")
    doc.add_picture(io.BytesIO(jpeg()))
    result = parse_docx(save(doc, tmp_path), tmp_path)
    assert result["assets"] == []
    assert result["skipped_object_count"] == 2
    assert "Figure 7: ambiguous" in render_word(result)
    assert not list((tmp_path / "images").glob("*"))


def test_numbered_captions_resolve_grouped_pictures(tmp_path):
    doc = Document()
    doc.add_picture(io.BytesIO(jpeg()))
    doc.add_picture(io.BytesIO(jpeg()))
    doc.add_paragraph("Figure 1: First", style="Caption")
    doc.add_paragraph("Figure 2: Second", style="Caption")
    result = parse_docx(save(doc, tmp_path))
    assert [a["caption"] for a in result["assets"]] == ["Figure 1: First", "Figure 2: Second"]


def test_headers_footers_textboxes_and_final_revision_view(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("Current")
    for kind, value in (("del", "Deleted"), ("ins", "Inserted"), ("moveFrom", "Moved old"), ("moveTo", "Moved new")):
        wrapper = OxmlElement("w:"+kind)
        run, text = OxmlElement("w:r"), OxmlElement("w:t")
        text.text = value
        run.append(text)
        wrapper.append(run)
        paragraph._p.append(wrapper)
    textbox = OxmlElement("w:txbxContent")
    text_p = doc.add_paragraph("Textbox data")
    textbox.append(text_p._p)
    paragraph._p.append(textbox)
    doc.sections[0].header.paragraphs[0].text = "Header data"
    doc.sections[0].footer.paragraphs[0].text = "Footer data"
    text = render_word(parse_docx(save(doc, tmp_path)))
    assert all(v in text for v in ("Current", "Inserted", "Moved new", "Textbox data", "Header data", "Footer data"))
    assert "Deleted" not in text and "Moved old" not in text


def test_note_parts_are_read_at_the_reference(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Before")
    ref = OxmlElement("w:footnoteReference")
    ref.set(f'{{{NS["w"]}}}id', "1")
    p.add_run()._r.append(ref)
    p.add_run("After")
    path = save(doc, tmp_path)
    with ZipFile(path) as archive:
        contents = {name: archive.read(name) for name in archive.namelist()}
    contents['word/footnotes.xml'] = f'<w:footnotes xmlns:w="{NS["w"]}"><w:footnote w:id="1"><w:p><w:r><w:t>Footnote data</w:t></w:r></w:p></w:footnote></w:footnotes>'.encode()
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, value in contents.items():
            archive.writestr(name, value)
    text = render_word(parse_docx(path))
    assert text.index("Before") < text.index("Footnote data") < text.index("After")


def test_table_caption_and_note_are_attached(tmp_path):
    doc = Document()
    doc.add_paragraph("Table 1: Pump limits", style="Caption")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Model"
    table.cell(0, 1).text = "Maximum"
    table.cell(1, 0).text = "ZQ-25"
    table.cell(1, 1).text = "100"
    table.cell(1, 1).add_table(rows=1, cols=1).cell(0, 0).text = "Nested specification"
    doc.add_paragraph("Note: values in kPa")
    result = parse_docx(save(doc, tmp_path))
    asset = result["assets"][0]
    assert asset["caption"] == "Table 1: Pump limits"
    assert asset["table"]["footnotes"] == ["Note: values in kPa"]


def test_caption_does_not_cross_section_boundary(tmp_path):
    doc = Document()
    doc.add_heading("First section", 1)
    doc.add_picture(io.BytesIO(jpeg()))
    doc.add_heading("Second section", 1)
    doc.add_paragraph("Figure 1: unrelated", style="Caption")
    result = parse_docx(save(doc, tmp_path))
    assert result["assets"][0]["caption"] == ""
    assert result["skipped_object_count"] == 0
    assert result["diagnostics"]


def test_missing_image_relationship_fails(tmp_path):
    doc = Document()
    shape = doc.add_picture(io.BytesIO(jpeg()))
    blip = next(shape._inline.iter(f'{{{NS["a"]}}}blip'))
    blip.set(f'{{{NS["r"]}}}embed', "missing")
    with pytest.raises(ValueError, match="Missing or external"):
        parse_docx(save(doc, tmp_path))


@pytest.mark.parametrize("strategy", ["character", "token", "sentence", "recursive", "structure"])
def test_word_chunk_strategies_preserve_table_headers(tmp_path, strategy):
    from rag_mcp.word_assets import parse_word
    from rag_mcp.chunk_strategies import ChunkingConfig

    doc = Document()
    doc.add_heading("Specification", 1)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Model", "Limit"
    table.cell(1, 0).text, table.cell(1, 1).text = "ZX256", "90"
    parsed = parse_word(save(doc, tmp_path), source_root=tmp_path, config=ChunkingConfig(strategy, 256, 16),
                        tokenizer_model="BAAI/bge-m3")
    rows = [c.text for c in parsed.chunks if c.metadata.get("asset_kind") == "table"]
    assert len(rows) == 2 and all("Model | Limit" in row for row in rows)
    assert "ZX256 | 90" in rows[-1]


def test_word_asset_and_index_rollback_on_publication_failure(tmp_path, monkeypatch):
    from types import SimpleNamespace
    from qdrant_client import QdrantClient
    from rag_mcp.config import RagConfig
    from rag_mcp.service_knowledge import DenseDocumentIngester, KnowledgeDocumentStore, KnowledgePaths
    from rag_mcp.storage.qdrant_store import QdrantStore
    from rag_mcp.sparse import Bm25SparseEncoder
    from rag_mcp.word_assets import read_word_content

    documents = KnowledgeDocumentStore(KnowledgePaths(tmp_path / "tenant"))
    client = QdrantClient(":memory:")
    store = QdrantStore(collection="word_rollback", client=client, namespace="tenant")
    config = RagConfig(backend="qdrant", collection="word_rollback")
    embedder = SimpleNamespace(embed_documents=lambda texts: [[1.0]+[0.0]*1023 for _ in texts])
    ingester = DenseDocumentIngester(config, embedder, Bm25SparseEncoder(), store)
    def content(caption):
        doc = Document()
        doc.add_paragraph("Figure 1: " + caption, style="Caption")
        doc.add_picture(io.BytesIO(jpeg()))
        stream = io.BytesIO()
        doc.save(stream)
        return stream.getvalue()
    try:
        original = content("Previous")
        added = documents.store_and_index("image.docx", io.BytesIO(original), ingester)
        source = documents.paths.knowledge_dir / "image.docx"
        before = read_word_content(source, documents.paths.knowledge_dir)
        def fail(_value):
            raise OSError("controlled disk failure")
        monkeypatch.setattr(documents, "_write_manifest", fail)
        with pytest.raises(OSError, match="controlled disk failure"):
            documents.store_and_index("image.docx", io.BytesIO(content("Replacement")), ingester)
        assert source.read_bytes() == original
        assert read_word_content(source, documents.paths.knowledge_dir) == before
        images = list((documents.paths.knowledge_dir / ".word-assets" / added["document_id"] / "images").glob("*"))
        assert len(images) == 1 and images[0].read_bytes() == jpeg()
        hits = store.dense_search([1.0]+[0.0]*1023, limit=10, filters={})
        assert any("Previous" in hit.payload["text"] for hit in hits)
        assert all("Replacement" not in hit.payload["text"] for hit in hits)
    finally:
        client.close()


def test_previous_word_parser_is_not_deduplicated(tmp_path):
    from test_mcp_http import ProtocolRuntime
    from rag_mcp.chunk_strategies import ChunkingConfig
    runtime = ProtocolRuntime(tmp_path)
    runtime.create_knowledge_base("tenant", ChunkingConfig(), runtime.config.embedding_model, runtime.config.reranker_model, **runtime.provider_settings)
    documents = runtime._tenant("tenant").documents
    (documents.paths.knowledge_dir / "old.docx").write_bytes(b"old word")
    documents._write_manifest({"old.docx": {"name": "old.docx", "sha256": "a"*64,
                               "chunking": ChunkingConfig().metadata("BAAI/bge-m3")}})
    result = runtime.check_documents("tenant", [{"name": "old.docx", "sha256": "a"*64}])
    assert result["files"][0]["document"] is None


def test_word_resource_pages_keep_inline_formulas_whole():
    from rag_mcp.resources import paragraph_pages
    equation = "$" + "x+" * 100 + "1$"
    text = "a" * 7950 + equation + "z" * 200
    pages = paragraph_pages(text, protect_inline_math=True)
    assert "".join(pages) == text
    assert sum(equation in page for page in pages) == 1
