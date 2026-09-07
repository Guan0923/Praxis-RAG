"""Independent source text/OOXML inspection and PDF page rendering for ground truth."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from PIL import Image, ImageDraw
from pypdf import PdfReader

from prepare import DEFAULT_OUTPUT


def inspect(output: Path) -> None:
    import pypdfium2 as pdfium

    destination = output / "source-review"
    destination.mkdir(parents=True, exist_ok=True)
    records = []
    for source in sorted((output / "sources").iterdir()):
        if source.suffix == ".pdf":
            reader = PdfReader(source)
            pages = [{"page": i + 1, "text": page.extract_text() or ""} for i, page in enumerate(reader.pages)]
            record = {"file": source.name, "pages": pages}
            thumbnails = []
            pdf = pdfium.PdfDocument(source)
            try:
                for i in range(len(pdf)):
                    page = pdf[i]
                    bitmap = page.render(scale=1.4)
                    image = bitmap.to_pil().convert("RGB")
                    image.save(destination / f"{source.stem}-page-{i + 1}.png")
                    image.thumbnail((380, 520))
                    tile = Image.new("RGB", (400, 550), "white")
                    tile.paste(image, ((400 - image.width) // 2, 22))
                    ImageDraw.Draw(tile).text((8, 5), f"{source.stem} / {i + 1}", fill="black")
                    thumbnails.append(tile)
                    bitmap.close()
                    page.close()
            finally:
                pdf.close()
            for start in range(0, len(thumbnails), 9):
                section = thumbnails[start:start + 9]
                sheet = Image.new("RGB", (1200, 550 * ((len(section) + 2) // 3)), "#dddddd")
                for i, tile in enumerate(section):
                    sheet.paste(tile, ((i % 3) * 400, (i // 3) * 550))
                sheet.save(destination / f"{source.stem}-sheet-{start // 9 + 1}.png")
            summary = "\n\n".join(f"PAGE {p['page']}\n{p['text']}" for p in pages)
        elif source.suffix == ".docx":
            document = Document(source)
            with zipfile.ZipFile(source) as archive:
                media = [
                    {"name": item.filename, "bytes": item.file_size}
                    for item in archive.infolist() if item.filename.startswith("word/media/")
                ]
                document_xml = archive.read("word/document.xml")
            tables = []
            for table_index, element in enumerate(document.element.body.iter(qn("w:tbl"))):
                table = Table(element, document._body)
                rows = []
                for row in table._tbl.tr_lst:
                    cells = []
                    for cell in row.tc_lst:
                        properties = cell.tcPr
                        cells.append({
                            "paragraphs": ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in cell.findall(qn("w:p"))],
                            "grid_span": properties.gridSpan.val if properties is not None and properties.gridSpan is not None else 1,
                            "vertical_merge": properties.vMerge.val if properties is not None and properties.vMerge is not None else None,
                        })
                    rows.append(cells)
                tables.append({"table": table_index, "rows": rows})
            paragraphs = [p.text for p in document.paragraphs]
            record = {
                "file": source.name,
                "paragraphs": paragraphs,
                "tables": tables,
                "media": media,
                "inline_drawings": len(document.inline_shapes),
                "math_objects": document_xml.count(b"<m:oMath>"),
                "math_paragraphs": document_xml.count(b"<m:oMathPara>"),
            }
            summary = (
                "\n".join(paragraphs)
                + "\n\nSTRUCTURE\n"
                + json.dumps({key: record[key] for key in ("media", "inline_drawings", "math_objects", "math_paragraphs")}, ensure_ascii=False, indent=2)
                + "\n\nTABLES\n"
                + json.dumps(tables, ensure_ascii=False, indent=2)
            )
        else:
            continue
        (destination / f"{source.stem}.txt").write_text(summary, encoding="utf-8")
        records.append(record)
        print(f"Reviewed {source.name}", flush=True)
    (destination / "source-content.json").write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    inspect(parser.parse_args().output.resolve())
